"""
Genera presentacion/RFKAN_Presentacion.pptx (editable: tablas y graficos nativos de
Office, notas del orador) a partir de content.py.
  lab/.venv/bin/python presentacion/build_pptx.py
"""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

import content as C

OUT = Path(__file__).parent / "RFKAN_Presentacion.pptx"
DARK, LIGHT, INK, SOFT = RGBColor(0x14, 0x21, 0x3D), RGBColor(0xF4, 0xF6, 0xF3), RGBColor(0x1F, 0x29, 0x33), RGBColor(0x4A, 0x55, 0x68)
TEAL, ORANGE, RULE = RGBColor(0x1B, 0x7F, 0x8C), RGBColor(0xD9, 0x6A, 0x1F), RGBColor(0xD5, 0xDB, 0xD6)
SERIES = [RGBColor(0x8A, 0x94, 0xA6), TEAL, RGBColor(0xE8, 0xB0, 0x7A), ORANGE,   # KAN RKAN FKAN RFKAN
          RGBColor(0x33, 0x33, 0x33), RGBColor(0x66, 0x66, 0x66), RGBColor(0x99, 0x99, 0x99)]   # LSTM GRU TCN
HEAD, BODY = "Source Serif 4", "IBM Plex Sans"
W, H = Inches(13.333), Inches(7.5)
X0, CW = Inches(0.7), Inches(11.93)


def txt(slide, x, y, w, h, text, size=18, color=INK, bold=False, font=BODY, align=None):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    if align:
        p.alignment = align
    for run in p.runs:
        run.font.size, run.font.bold, run.font.name = Pt(size), bold, font
        run.font.color.rgb = color
    return tb


def bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def table(slide, x, y, w, header, rows, size=13):
    t = slide.shapes.add_table(len(rows) + 1, len(header), x, y, w, Inches(0.36) * (len(rows) + 1)).table
    for i, h in enumerate(header):
        c = t.cell(0, i)
        c.text = h
        c.fill.solid()
        c.fill.fore_color.rgb = DARK
    for k, row in enumerate(rows):
        for i, v in enumerate(row):
            c = t.cell(k + 1, i)
            c.text = str(v)
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if k % 2 == 0 else RGBColor(0xEC, 0xF0, 0xEE)
    for r_ in t.rows:
        for c in r_.cells:
            for p in c.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size, run.font.name = Pt(size), BODY
    for c in t.rows[0].cells:
        for run in c.text_frame.paragraphs[0].runs:
            run.font.bold, run.font.color.rgb = True, LIGHT
    return Inches(0.36) * (len(rows) + 1)


def image(slide, name, x, y, w, h):
    path = C.ANA / name
    if not path.exists():
        txt(slide, x, y, w, Inches(0.5), f"[figura pendiente: {name}]", 14, SOFT)
        return Inches(0.5)
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = int(iw * scale), int(ih * scale)
    slide.shapes.add_picture(str(path), x + (w - pw) // 2, y, pw, ph)
    return ph


def render(prs, s, n):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if s["id"] == "cover":
        bg(slide, DARK)
        txt(slide, X0, Inches(2.4), CW, Inches(1.4), s["title"], 48, LIGHT, True, HEAD)
        txt(slide, X0, Inches(3.9), Inches(10), Inches(1.2), s["body"][0][1], 20, RGBColor(0xBF, 0xD8, 0xD5))
        slide.notes_slide.notes_text_frame.text = s["notes"]
        return
    bg(slide, LIGHT)
    txt(slide, X0, Inches(0.25), CW, Inches(0.3), C.LABEL[s["section"]].upper(), 12, TEAL, True)
    txt(slide, X0, Inches(0.5), CW, Inches(0.9), s["title"], 32, DARK, True, HEAD)
    txt(slide, Inches(12.2), Inches(7.0), Inches(0.8), Inches(0.3), str(n), 11, SOFT, align=PP_ALIGN.RIGHT)
    y = Inches(1.55)
    kinds = [b[0] for b in s["body"]]
    side = "table" in kinds and "image" in kinds
    for b in s["body"]:
        k = b[0]
        if k == "text":
            txt(slide, X0, y, CW, Inches(0.9), b[1], 17, SOFT)
            y += Inches(0.95)
        elif k == "bullets":
            tb = slide.shapes.add_textbox(X0, y, CW, Inches(0.5) * len(b[1]))
            tf = tb.text_frame
            tf.word_wrap = True
            for i, item in enumerate(b[1]):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = "•  " + item
                p.space_after = Pt(8)
                for run in p.runs:
                    run.font.size, run.font.name, run.font.color.rgb = Pt(19), BODY, INK
            y += Inches(0.52) * len(b[1])
        elif k == "eq":
            for line in b[1]:
                txt(slide, X0, y, CW, Inches(0.6), line, 24, DARK, font="Cambria Math", align=PP_ALIGN.CENTER)
                y += Inches(0.65)
            y += Inches(0.15)
        elif k == "big":
            w = CW // len(b[1])
            for i, (val, lab) in enumerate(b[1]):
                txt(slide, X0 + w * i, y, w, Inches(0.9), val, 40, ORANGE, True, HEAD)
                txt(slide, X0 + w * i, y + Inches(0.85), w - Inches(0.2), Inches(0.6), lab, 15, SOFT)
            y += Inches(1.7)
        elif k == "table":
            tw = Inches(6.2) if side else CW
            size = 13 if len(b[2]) <= 6 else 12
            h = table(slide, X0, y, tw, b[1], b[2], size)
            if not side:
                y += h + Inches(0.25)
        elif k == "image":
            if side:
                image(slide, b[1], X0 + Inches(6.5), y, Inches(5.4), Inches(4.6))
                y += Inches(4.7)
            else:
                y += image(slide, b[1], X0, y, CW, Inches(7.0) - y - Inches(1.0)) + Inches(0.1)
        elif k == "chart_line":
            cd = CategoryChartData()
            cd.categories = b[1]
            for name, vals in b[2].items():
                cd.add_series(name, [None if v != v else round(v, 4) for v in vals])
            ch = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, X0, y, CW, Inches(4.2), cd).chart
            ch.has_legend, ch.legend.position, ch.legend.include_in_layout = True, XL_LEGEND_POSITION.RIGHT, False
            ch.value_axis.has_title = True
            ch.value_axis.axis_title.text_frame.text = b[3]
            for i, (ser, col) in enumerate(zip(ch.series, SERIES)):
                ser.format.line.color.rgb = col
                ser.format.line.width = Pt(2.5 if i < 4 else 1.5)
                if i >= 4:
                    ser.format.line.dash_style = 4   # MSO_LINE.DASH
                ser.marker.format.fill.solid()
                ser.marker.format.fill.fore_color.rgb = col
            y += Inches(4.3)
        elif k == "flow":
            n_ = len(b[1])
            gap = Inches(0.25)
            w = (CW - gap * (n_ - 1)) // n_
            for i, lab in enumerate(b[1]):
                box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, X0 + (w + gap) * i, y, w, Inches(1.1))
                box.fill.solid()
                box.fill.fore_color.rgb = TEAL if "Capa" in lab else (ORANGE if "ŷ" in lab else RGBColor(0xFF, 0xFF, 0xFF))
                box.line.color.rgb = RULE
                tf = box.text_frame
                tf.word_wrap = True
                tf.text = lab
                for run in tf.paragraphs[0].runs:
                    run.font.size, run.font.name = Pt(13), BODY
                    run.font.color.rgb = LIGHT if ("Capa" in lab or "ŷ" in lab) else INK
                if i < n_ - 1:
                    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, X0 + (w + gap) * i + w + Inches(0.03),
                                               y + Inches(0.45), gap - Inches(0.06), Inches(0.2))
                    a.fill.solid()
                    a.fill.fore_color.rgb = SOFT
                    a.line.fill.background()
            y += Inches(1.45)
    slide.notes_slide.notes_text_frame.text = s["notes"]


prs = Presentation()
prs.slide_width, prs.slide_height = W, H
for n, s in enumerate(C.SLIDES, 1):
    render(prs, s, n)
prs.save(OUT)
print("->", OUT, len(C.SLIDES), "slides")

# guion del video: las mismas notas del orador, por bloque
lines = ["# Guion del video (≤ 18 min: Metodología 6 / Implementación 6 / Resultados 6)", "",
         "Generado por build_pptx.py desde las notas de content.py.", ""]
sec = None
for n, s in enumerate(C.SLIDES, 1):
    if s["section"] != sec:
        sec = s["section"]
        lines += [f"## {C.LABEL[sec]}", ""]
    lines += [f"**{n}. {s['title']}**", "", s["notes"], ""]
(Path(__file__).parent / "SPEECH.md").write_text("\n".join(lines))
print("->", Path(__file__).parent / "SPEECH.md")
